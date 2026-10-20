"""Informes: un único modelo de documento que se escribe en Word (.docx), PDF (HTML → Edge/Chrome) y Excel (.xlsx).

Tipos:
  · sesion       – informe individual de una sesión (para el jugador o el cuerpo técnico, lenguaje sencillo).
  · tecnico      – informe técnico de una sesión: todas las variables, zonas, minuto a minuto, WCS, eventos, calidad
                   de la señal, canales del archivo y parámetros de procesado.
  · jugador      – evolución de un jugador en las sesiones seleccionadas.
  · colectivo    – comparación de los jugadores en las sesiones seleccionadas (p. ej., un entrenamiento).
  · estadistico  – informe técnico colectivo: descriptivo, comparaciones, correlaciones, variabilidad entre sesiones y
                   concordancia con SPRO, con el apartado de métodos.
"""
from __future__ import annotations

import base64
import html
import math
import re
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from . import analisis as A
from . import config as C
from . import estadistica as E
from . import figuras as F
from . import metricas as M
from . import servicio as SV
from . import spro as SP
from .pdf import html_a_pdf

APP = "WIMU Lab"


# =========================================================================== formato
def fnum(x, d=0):
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "—"
    try:
        s = f"{float(x):,.{d}f}"
    except (TypeError, ValueError):
        return str(x)
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def fvar(k, x):
    v = M.VAR.get(k)
    return fnum(x, v[3] if v else 1)


def ffecha(iso, hora=True):
    if not iso:
        return "—"
    try:
        d = datetime.fromisoformat(str(iso))
        return d.strftime("%d/%m/%Y %H:%M") if hora else d.strftime("%d/%m/%Y")
    except ValueError:
        return str(iso)


def fmin(seg):
    if seg is None:
        return "—"
    seg = int(round(seg))
    return f"{seg // 60}:{seg % 60:02d}"


def _slug(s):
    s = re.sub(r"[^\w\-]+", "_", str(s), flags=re.UNICODE).strip("_")
    return s[:60] or "informe"


# =========================================================================== modelo de documento
class Doc:
    def __init__(self, titulo, subtitulo="", meta=None):
        self.titulo, self.subtitulo, self.meta = titulo, subtitulo, meta or []
        self.b = []

    def h1(self, t):
        self.b.append(("h1", t))

    def h2(self, t):
        self.b.append(("h2", t))

    def p(self, t):
        self.b.append(("p", t))

    def nota(self, t):
        self.b.append(("nota", t))

    def lista(self, items):
        self.b.append(("lista", list(items)))

    def kpis(self, items):
        self.b.append(("kpis", [(a, b, c) for a, b, c in items]))

    def tabla(self, cols, filas, nota=None, hoja=None, num=None, datos=None):
        """filas: textos ya formateados; datos (opcional): mismos valores sin formatear para Excel."""
        self.b.append(("tabla", dict(cols=cols, filas=filas, nota=nota, hoja=hoja, datos=datos)))

    def img(self, ruta, pie="", ancho=1.0):
        if ruta:
            self.b.append(("img", (ruta, pie, ancho)))

    def salto(self):
        self.b.append(("salto", None))


# --------------------------------------------------------------------------- Word
def a_docx(doc: Doc, destino: Path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    navy = RGBColor(0x14, 0x2B, 0x4C)

    def sombrear(celda, color):
        tcPr = celda._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    p = d.add_paragraph()
    r = p.add_run(doc.titulo)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), navy
    if doc.subtitulo:
        p = d.add_paragraph()
        r = p.add_run(doc.subtitulo)
        r.font.size, r.font.color.rgb = Pt(11), RGBColor(0x5B, 0x65, 0x77)
    if doc.meta:
        p = d.add_paragraph()
        r = p.add_run("  ·  ".join(doc.meta))
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor(0x5B, 0x65, 0x77)
    for tipo, x in doc.b:
        if tipo == "h1":
            h = d.add_heading(x, level=1)
            for r in h.runs:
                r.font.color.rgb = navy
                r.font.size = Pt(13)
        elif tipo == "h2":
            h = d.add_heading(x, level=2)
            for r in h.runs:
                r.font.color.rgb = navy
                r.font.size = Pt(11)
        elif tipo == "p":
            d.add_paragraph(x)
        elif tipo == "nota":
            p = d.add_paragraph()
            r = p.add_run(x)
            r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0x5B, 0x65, 0x77)
        elif tipo == "lista":
            for it in x:
                d.add_paragraph(it, style="List Bullet")
        elif tipo == "kpis":
            ncol = 4
            filas = math.ceil(len(x) / ncol)
            t = d.add_table(rows=filas * 2, cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (lab, val, uni) in enumerate(x):
                f, c = divmod(i, ncol)
                cv, cl = t.cell(2 * f, c), t.cell(2 * f + 1, c)
                for cc in (cv, cl):
                    sombrear(cc, "EEF3FA")
                pv = cv.paragraphs[0]
                pv.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = pv.add_run(str(val))
                r.bold, r.font.size, r.font.color.rgb = True, Pt(15), navy
                if uni:
                    r2 = pv.add_run(f" {uni}")
                    r2.font.size = Pt(8.5)
                pl = cl.paragraphs[0]
                pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = pl.add_run(lab)
                r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x5B, 0x65, 0x77)
            d.add_paragraph()
        elif tipo == "tabla":
            cols, filas = x["cols"], x["filas"]
            t = d.add_table(rows=1 + len(filas), cols=len(cols))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, c in enumerate(cols):
                cell = t.cell(0, j)
                sombrear(cell, "142B4C")
                r = cell.paragraphs[0].add_run(str(c))
                r.bold, r.font.size, r.font.color.rgb = True, Pt(8), RGBColor(0xFF, 0xFF, 0xFF)
            for i, fila in enumerate(filas):
                for j, v in enumerate(fila):
                    cell = t.cell(i + 1, j)
                    if i % 2:
                        sombrear(cell, "F4F6FA")
                    par = cell.paragraphs[0]
                    if j > 0:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if re.fullmatch(r"[\s\d.,−\-–+%<>—:()\[\]]*", str(v)) else WD_ALIGN_PARAGRAPH.LEFT
                    r = par.add_run(str(v))
                    r.font.size = Pt(8)
            if x.get("nota"):
                p = d.add_paragraph()
                r = p.add_run(x["nota"])
                r.italic, r.font.size, r.font.color.rgb = True, Pt(8), RGBColor(0x5B, 0x65, 0x77)
            else:
                d.add_paragraph()
        elif tipo == "img":
            ruta, pie, ancho = x
            d.add_picture(ruta, width=Cm(16.5 * ancho))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if pie:
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(pie)
                r.italic, r.font.size = True, Pt(8.5)
        elif tipo == "salto":
            d.add_page_break()
    pie = sec.footer.paragraphs[0]
    pie.text = f"{APP} · {datetime.now():%d/%m/%Y %H:%M}"
    pie.runs[0].font.size = Pt(8)
    destino.parent.mkdir(parents=True, exist_ok=True)
    d.save(destino)
    return destino


# --------------------------------------------------------------------------- HTML / PDF
CSS = """
@page { size: A4; margin: 16mm 15mm 16mm 15mm; }
* { box-sizing: border-box; }
body { font-family: "Segoe UI", Calibri, Arial, sans-serif; color: #1f2733; font-size: 10pt; line-height: 1.42; margin: 0; }
.cab { border-bottom: 3px solid #142B4C; padding-bottom: 8px; margin-bottom: 14px; }
.cab h1 { font-size: 20pt; color: #142B4C; margin: 0 0 2px; }
.cab .sub { color: #5b6577; font-size: 11pt; }
.cab .meta { color: #5b6577; font-size: 8.5pt; margin-top: 4px; }
h2 { color: #142B4C; font-size: 13pt; margin: 18px 0 6px; border-bottom: 1px solid #d8dde6; padding-bottom: 3px; break-after: avoid; }
h3 { color: #142B4C; font-size: 11pt; margin: 12px 0 4px; break-after: avoid; }
p { margin: 4px 0 8px; }
.nota { color: #5b6577; font-size: 8.5pt; font-style: italic; }
.kpis { display: grid; grid-template-columns: repeat(4, 1fr); gap: 6px; margin: 8px 0 12px; break-inside: avoid; }
.kpi { background: #eef3fa; border-radius: 6px; padding: 7px 6px; text-align: center; }
.kpi b { display: block; font-size: 15pt; color: #142B4C; }
.kpi b small { font-size: 8pt; font-weight: 400; color: #5b6577; }
.kpi span { font-size: 8pt; color: #5b6577; }
table { border-collapse: collapse; width: 100%; font-size: 8pt; margin: 4px 0 2px; break-inside: auto; }
th { background: #142B4C; color: #fff; text-align: left; padding: 4px 5px; font-weight: 600; }
td { padding: 3px 5px; border-bottom: 1px solid #e6e9ef; }
td.n { text-align: right; font-variant-numeric: tabular-nums; }
tr:nth-child(even) td { background: #f6f8fb; }
tr { break-inside: avoid; }
figure { margin: 8px 0 12px; text-align: center; break-inside: avoid; }
figure img { max-width: 100%; }
figcaption { font-size: 8.5pt; font-style: italic; color: #5b6577; }
.salto { break-after: page; }
ul { margin: 4px 0 8px 18px; padding: 0; }
.pie { position: fixed; bottom: -10mm; left: 0; right: 0; font-size: 7.5pt; color: #8a94a6; text-align: right; }
"""


def a_html(doc: Doc, destino: Path):
    e = html.escape
    out = [f"<!doctype html><html lang='es'><head><meta charset='utf-8'><title>{e(doc.titulo)}</title><style>{CSS}</style></head><body>",
           f"<div class='cab'><h1>{e(doc.titulo)}</h1>" + (f"<div class='sub'>{e(doc.subtitulo)}</div>" if doc.subtitulo else "")
           + (f"<div class='meta'>{e('  ·  '.join(doc.meta))}</div>" if doc.meta else "") + "</div>"]
    for tipo, x in doc.b:
        if tipo == "h1":
            out.append(f"<h2>{e(x)}</h2>")
        elif tipo == "h2":
            out.append(f"<h3>{e(x)}</h3>")
        elif tipo == "p":
            out.append(f"<p>{e(x)}</p>")
        elif tipo == "nota":
            out.append(f"<p class='nota'>{e(x)}</p>")
        elif tipo == "lista":
            out.append("<ul>" + "".join(f"<li>{e(i)}</li>" for i in x) + "</ul>")
        elif tipo == "kpis":
            out.append("<div class='kpis'>" + "".join(f"<div class='kpi'><b>{e(str(v))}" + (f" <small>{e(u)}</small>" if u else "")
                                                       + f"</b><span>{e(l)}</span></div>" for l, v, u in x) + "</div>")
        elif tipo == "tabla":
            num = lambda v: bool(re.fullmatch(r"[\s\d.,−\-–+%<>—:()\[\]]*", str(v)))  # noqa: E731
            out.append("<table><thead><tr>" + "".join(f"<th>{e(str(c))}</th>" for c in x["cols"]) + "</tr></thead><tbody>"
                       + "".join("<tr>" + "".join(f"<td class='{'n' if j and num(v) else ''}'>{e(str(v))}</td>" for j, v in enumerate(f)) + "</tr>"
                                 for f in x["filas"]) + "</tbody></table>")
            if x.get("nota"):
                out.append(f"<p class='nota'>{e(x['nota'])}</p>")
        elif tipo == "img":
            ruta, pie, ancho = x
            b64 = base64.b64encode(Path(ruta).read_bytes()).decode()
            out.append(f"<figure><img style='width:{ancho * 100:.0f}%' src='data:image/png;base64,{b64}'>" + (f"<figcaption>{e(pie)}</figcaption>" if pie else "") + "</figure>")
        elif tipo == "salto":
            out.append("<div class='salto'></div>")
    out.append(f"<div class='pie'>{APP} · {datetime.now():%d/%m/%Y %H:%M}</div></body></html>")
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text("".join(out), encoding="utf-8")
    return destino


# --------------------------------------------------------------------------- Excel
def a_xlsx(doc: Doc, destino: Path, extra: dict | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([doc.titulo])
    ws["A1"].font = Font(bold=True, size=14, color="142B4C")
    if doc.subtitulo:
        ws.append([doc.subtitulo])
    for m in doc.meta:
        ws.append([m])
    ws.append([])
    for tipo, x in doc.b:
        if tipo == "kpis":
            for l, v, u in x:
                ws.append([l, _a_num(v), u])
    hojas = {}
    cab = PatternFill("solid", fgColor="142B4C")

    def hoja(nombre, cols, filas):
        nombre = re.sub(r"[\[\]\*\?/\\:]", "", nombre)[:31] or "Datos"
        base, k = nombre, 2
        while nombre in hojas:
            nombre = f"{base[:28]} {k}"
            k += 1
        w = wb.create_sheet(nombre)
        hojas[nombre] = w
        w.append(list(cols))
        for c in w[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), cab, Alignment(wrap_text=True, vertical="center")
        for f in filas:
            w.append([_a_num(v) for v in f])
        for j, c in enumerate(cols, 1):
            largo = max([len(str(c))] + [len(str(f[j - 1])) for f in filas[:300] if j - 1 < len(f)])
            w.column_dimensions[get_column_letter(j)].width = min(48, max(9, largo + 2))
        w.freeze_panes = "B2"

    for tipo, x in doc.b:
        if tipo == "tabla" and x.get("hoja"):
            hoja(x["hoja"], x["cols"], x["datos"] if x.get("datos") is not None else x["filas"])
    for nombre, df in (extra or {}).items():
        if df is not None and len(df):
            hoja(nombre, list(df.columns), df.astype(object).where(df.notna(), None).values.tolist())
    ws.column_dimensions["A"].width = 44
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def _a_num(v):
    if isinstance(v, (int, float, np.integer, np.floating)):
        return None if isinstance(v, float) and not math.isfinite(v) else float(v) if isinstance(v, (float, np.floating)) else int(v)
    if isinstance(v, str):
        s = v.strip().replace(".", "").replace(",", ".").replace("−", "-")
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            return float(s)
        if v.strip() == "—":
            return None
    return v


# =========================================================================== textos auxiliares
KPI_SESION = ["duracion_min", "distancia_m", "m_min", "v_max_kmh", "hsr_m", "n_sprints", "n_acel", "n_decel", "player_load", "fc_media",
              "fc_max", "trimp"]
KPI_LABEL = {"duracion_min": ("Duración", "min"), "distancia_m": ("Distancia", "m"), "m_min": ("Distancia relativa", "m/min"),
             "v_max_kmh": ("Velocidad máxima", "km/h"), "hsr_m": ("Alta velocidad", "m"), "n_sprints": ("Sprints", ""),
             "n_acel": ("Aceleraciones", ""), "n_decel": ("Deceleraciones", ""), "player_load": ("Player Load", "u.a."),
             "fc_media": ("FC media", "lpm"), "fc_max": ("FC máxima", "lpm"), "trimp": ("TRIMP", "u.a."), "pct_fcmax_media": ("FC media", "% FCmáx")}
TABLA_JUG = ["duracion_min", "distancia_m", "m_min", "v_max_kmh", "hsr_m", "sprint_m", "n_sprints", "n_acel", "n_decel", "player_load",
             "pl_min", "fc_media", "fc_max", "pct_fcmax_media", "trimp"]


def _kpis(res, claves=KPI_SESION):
    out = []
    for k in claves:
        if res.get(k) is None:
            continue
        lab, u = KPI_LABEL.get(k, (M.etiqueta(k), ""))
        out.append((lab, fvar(k, res[k]), u))
    return out


def _parametros(cfg):
    zv = ", ".join(fnum(x, 0) for x in M.lista(cfg["zonas_velocidad"]))
    return [
        f"Distancia: {'integral de la velocidad Doppler del GPS' if cfg['fuente_distancia'] == 'velocidad' else 'canal DIST del dispositivo'}.",
        f"Muestras GPS válidas: ≥ {fnum(M._num(cfg['sat_min']), 0)} satélites, PDOP ≤ {fnum(M._num(cfg['pdop_max']), 1)} y velocidad ≤ {fnum(M._num(cfg['v_max_valida']), 0)} km/h; "
        f"huecos ≤ 1 s interpolados; velocidad suavizada con media móvil de {fnum(M._num(cfg['suavizado_v_s']), 1)} s.",
        f"Zonas de velocidad (km/h): {zv}. Alta velocidad (HSR) ≥ {fnum(M._num(cfg['umbral_hsr']), 0)} km/h; sprint ≥ {fnum(M._num(cfg['umbral_sprint']), 0)} km/h "
        f"durante ≥ {fnum(M._num(cfg['sprint_min_s']), 1)} s.",
        f"Aceleraciones y deceleraciones: |a| ≥ {fnum(M._num(cfg['umbral_acel']), 1)} m/s² (intensas ≥ {fnum(M._num(cfg['umbral_acel_alta']), 1)} m/s²) "
        f"durante ≥ {fnum(M._num(cfg['acel_min_s']), 1)} s; a = derivada de la velocidad suavizada a 1 s.",
        "Acelerómetro (100 Hz): AcelT = módulo de la aceleración (g, incluye la gravedad); Player Load = Σ√(Δx²+Δy²+Δz²)/100; "
        f"impactos = picos de AcelT por encima de {', '.join(fnum(x, 0) for x in M.lista(cfg['umbrales_impacto']))} g separados ≥ 0,25 s.",
        f"Frecuencia cardiaca: valores 30–230 lpm sin saltos > 25 lpm/s; zonas (% FCmáx): {', '.join(fnum(x, 0) for x in M.lista(cfg['zonas_fc']))}; "
        "TRIMP de Edwards (minutos en 50–60, 60–70, 70–80, 80–90 y ≥ 90 % FCmáx × 1 a 5).",
        f"Escenarios más exigentes (WCS): máximo de la distancia, HSR y Player Load en ventanas móviles de {', '.join(fnum(x, 0) for x in M.lista(cfg['ventanas_wcs']))} min.",
    ]


def _comparar_con_historial(sid, res, jid):
    """Frases comparando la sesión con la media del jugador en sus otras sesiones del mismo tipo."""
    if not jid:
        return []
    s = SV.reg().sesion(sid)
    otras = [x for x in SV.reg().sesiones(incluir_faltan=False) if x.get("jugador") == jid and x["id"] != sid and not x.get("excluida")
             and x.get("tipo") == s.get("tipo") and x.get("kpis")]
    if len(otras) < 2:
        return []
    frases = []
    for k, nombre in (("distancia_m", "distancia"), ("hsr_m", "distancia a alta velocidad"), ("player_load", "Player Load"), ("fc_media", "FC media")):
        vals = [o["kpis"].get(k) for o in otras if o["kpis"].get(k) is not None]
        if len(vals) < 2 or res.get(k) is None or not np.mean(vals):
            continue
        dif = 100 * (res[k] - np.mean(vals)) / np.mean(vals)
        if abs(dif) < 5:
            frases.append(f"{nombre.capitalize()}: similar a tu media ({fvar(k, np.mean(vals))}).")
        else:
            frases.append(f"{nombre.capitalize()}: un {fnum(abs(dif), 0)} % {'más' if dif > 0 else 'menos'} que tu media de {len(vals)} sesiones "
                          f"({fvar(k, np.mean(vals))}).")
    return frases


# =========================================================================== informes de sesión
def informe_sesion(sid, carpeta_fig: Path, tecnico=False):
    cfg = SV.cfg()
    a = SV.analisis_sesion(sid)
    s, r, meta = a["sesion"], a["resumen"], a["meta"]
    jug = s.get("jugador_nombre") or f"{s.get('dispositivo') or 'WIMU'} (jugador sin asignar)"
    titulo = ("Informe técnico de sesión" if tecnico else "Informe de sesión") + f" · {jug}"
    sub = " · ".join(x for x in (ffecha(meta.get("inicio")), s.get("tipo"), s.get("tarea")) if x)
    doc = Doc(titulo, sub, [f"Archivo: {s['archivo']}", f"Dispositivo: {meta.get('dispositivo') or '—'}",
                            f"Tiempo analizado: {fmin(a['desde'])}–{fmin(a['hasta'])} de {fmin(a['total'])}"])
    pref = carpeta_fig / f"{sid}_{'t' if tecnico else 'u'}"
    doc.h1("Resumen")
    doc.kpis(_kpis(r))
    frases = []
    if r.get("distancia_m") is not None:
        frases.append(f"Se recorrieron {fvar('distancia_m', r['distancia_m'])} m en {fnum(r['duracion_min'], 0)} min ({fvar('m_min', r['m_min'])} m/min), "
                      f"con una velocidad máxima de {fvar('v_max_kmh', r.get('v_max_kmh'))} km/h.")
        frases.append(f"A alta velocidad (≥ {fnum(M._num(cfg['umbral_hsr']), 0)} km/h) se recorrieron {fvar('hsr_m', r['hsr_m'])} m y se realizaron "
                      f"{r.get('n_sprints', 0)} sprints (≥ {fnum(M._num(cfg['umbral_sprint']), 0)} km/h), {r.get('n_acel', 0)} aceleraciones y "
                      f"{r.get('n_decel', 0)} deceleraciones (≥ {fnum(M._num(cfg['umbral_acel']), 0)} m/s²).")
    if r.get("fc_media") is not None:
        frases.append(f"La frecuencia cardiaca media fue {fvar('fc_media', r['fc_media'])} lpm ({fnum(r.get('pct_fcmax_media'), 0)} % de la FCmáx de "
                      f"{fnum(r.get('fcmax_ref'), 0)} lpm, {r.get('fcmax_fuente')}) y la máxima {fvar('fc_max', r['fc_max'])} lpm; "
                      f"{fnum(r.get('t_fc_alta_min'), 1)} min en zonas altas.")
    for f in frases:
        doc.p(f)
    hist = _comparar_con_historial(sid, r, s.get("jugador"))
    if hist:
        doc.h2("Comparación con tus sesiones anteriores" if not tecnico else "Comparación con la media del jugador (mismo tipo de sesión)")
        doc.lista(hist)
    for av in a["avisos"]:
        doc.nota("Aviso: " + av)
    doc.img(F.mapa(a["series"], f"{pref}_mapa.png"), "Recorrido en el terreno coloreado por la velocidad.", .72)
    doc.img(F.linea_tiempo(a["series"], a["eventos"], f"{pref}_tiempo.png", M._num(cfg["umbral_sprint"], 21), (a["desde"], a["hasta"])),
            "Velocidad (azul) y frecuencia cardiaca (rojo). Triángulos: sprints.")
    doc.h1("Intensidad: tiempo y distancia por zonas")
    if a["tablas"].get("velocidad"):
        doc.img(F.barras_zonas(a["tablas"]["velocidad"], "distancia_m", f"{pref}_zv.png", "Distancia por zonas de velocidad", "m"), ancho=.8)
    if a["tablas"].get("fc"):
        doc.img(F.barras_zonas(a["tablas"]["fc"], "tiempo_min", f"{pref}_zfc.png", "Tiempo por zonas de frecuencia cardiaca", "min",
                               ["#9cc3ee", "#5598e7", "#f2b33d", "#eb6834", "#d64545"]), ancho=.8)
    if a["segmentos"]:
        doc.h1("Partes de la sesión")
        ks = ["duracion_min", "distancia_m", "m_min", "v_max_kmh", "hsr_m", "n_acel", "n_decel", "player_load", "fc_media"]
        cols = ["Parte", "Tiempo"] + [M.etiqueta(k) for k in ks]
        filas = [[g["nombre"], f"{fmin(g['desde'])}–{fmin(g['hasta'])}"] + [fvar(k, g.get(k)) for k in ks] for g in a["segmentos"]]
        datos = [[g["nombre"], f"{fmin(g['desde'])}–{fmin(g['hasta'])}"] + [g.get(k) for k in ks] for g in a["segmentos"]]
        doc.tabla(cols, filas, hoja="Partes", datos=datos)
    if not tecnico:
        doc.h1("¿Qué significan estos datos?")
        doc.lista(["Distancia relativa (m/min): metros recorridos por cada minuto; permite comparar sesiones de distinta duración.",
                   f"Alta velocidad y sprints: carrera por encima de {fnum(M._num(cfg['umbral_hsr']), 0)} y {fnum(M._num(cfg['umbral_sprint']), 0)} km/h.",
                   "Aceleraciones y deceleraciones: arrancadas y frenadas intensas; las deceleraciones son las más exigentes para los músculos.",
                   "Player Load: suma de todos los movimientos registrados por el acelerómetro (saltos, contactos, cambios de dirección).",
                   "TRIMP: carga interna; combina el tiempo y la intensidad de la frecuencia cardiaca."])
        return doc, a
    # ------------------------------------------------------------------ técnico
    doc.salto()
    doc.h1("Todas las variables")
    filas, datos = [], []
    for k, nom, u, dec, grupo in M.VARIABLES:
        if r.get(k) is None:
            continue
        filas.append([grupo, nom, u, fnum(r[k], dec)])
        datos.append([grupo, nom, u, r[k]])
    for k in r:
        if k.startswith("impactos_") and k not in M.VAR:
            filas.append(["Carga externa · acelerómetro", M.etiqueta(k), "n", fnum(r[k], 0)])
            datos.append(["Carga externa · acelerómetro", M.etiqueta(k), "n", r[k]])
    doc.tabla(["Grupo", "Variable", "Unidad", "Valor"], filas, hoja="Variables", datos=datos,
              nota=f"FCmáx de referencia: {fnum(r.get('fcmax_ref'), 0)} lpm ({r.get('fcmax_fuente', '—')}).")
    nombres_z = dict(velocidad="Zonas de velocidad", aceleracion="Zonas de aceleración", acelt="Zonas de AcelT (acelerómetro)", fc="Zonas de frecuencia cardiaca")
    for clave, tit in nombres_z.items():
        t = a["tablas"].get(clave)
        if not t:
            continue
        doc.h2(tit)
        cols = [c for c in ("zona", "distancia_m", "pct_dist", "tiempo_min", "pct_tiempo") if c in t[0]]
        cab = dict(zona="Zona", distancia_m="Distancia (m)", pct_dist="% distancia", tiempo_min="Tiempo (min)", pct_tiempo="% tiempo")
        dec = dict(distancia_m=0, pct_dist=1, tiempo_min=2, pct_tiempo=1)
        doc.tabla([cab[c] for c in cols], [[f[c] if c == "zona" else fnum(f[c], dec[c]) for c in cols] for f in t], hoja=tit[:31],
                  datos=[[f[c] for c in cols] for f in t])
    if a["tablas"].get("aceleracion"):
        doc.img(F.barras_zonas(a["tablas"]["aceleracion"], "tiempo_min", f"{pref}_za.png", "Tiempo por zonas de aceleración", "min",
                               ["#8e2b2b", "#d64545", "#eb6834", "#c9d3e0", "#5598e7", "#2a78d6", "#142B4C"]), ancho=.8)
    if a["tablas"].get("acelt"):
        doc.img(F.barras_zonas(a["tablas"]["acelt"], "tiempo_min", f"{pref}_zt.png", "Tiempo por zonas de AcelT", "min"), ancho=.8)
    doc.h1("Escenarios más exigentes (WCS)")
    if a["wcs"]:
        cols = ["Ventana", "Distancia (m)", "m/min", "Inicio", "HSR (m)", "Inicio HSR", "Player Load", "Inicio PL"]
        filas = [[f"{fnum(w['ventana_min'], 0)} min", fnum(w.get("distancia"), 0), fnum(w.get("distancia_min"), 1), fmin(w.get("distancia_inicio_s")),
                  fnum(w.get("hsr"), 0), fmin(w.get("hsr_inicio_s")), fnum(w.get("player_load"), 1), fmin(w.get("player_load_inicio_s"))] for w in a["wcs"]]
        doc.tabla(cols, filas, hoja="WCS", datos=[[w["ventana_min"], w.get("distancia"), w.get("distancia_min"), w.get("distancia_inicio_s"), w.get("hsr"),
                                                  w.get("hsr_inicio_s"), w.get("player_load"), w.get("player_load_inicio_s")] for w in a["wcs"]],
                  nota="Máximo móvil con resolución de 1 s. «Inicio»: minuto:segundo desde el comienzo del registro.")
    doc.h1("Perfil minuto a minuto")
    doc.img(F.por_minuto(a["tablas"]["minutos"], f"{pref}_min.png"), "Distancia y HSR por minuto y FC media.")
    mins = a["tablas"]["minutos"]
    if mins:
        cl = [c for c in ("minuto", "distancia_m", "hsr_m", "v_max_kmh", "acel", "decel", "player_load", "fc_media") if c in mins[0]]
        cab = dict(minuto="Minuto", distancia_m="Distancia (m)", hsr_m="HSR (m)", v_max_kmh="V. máx. (km/h)", acel="Acel.", decel="Decel.",
                   player_load="Player Load", fc_media="FC media")
        dec = dict(minuto=0, distancia_m=0, hsr_m=0, v_max_kmh=1, acel=0, decel=0, player_load=1, fc_media=0)
        doc.tabla([cab[c] for c in cl], [[fnum(f.get(c), dec[c]) for c in cl] for f in mins], hoja="Minuto a minuto",
                  datos=[[f.get(c) for c in cl] for f in mins])
    if a["eventos"]:
        doc.h1("Esfuerzos registrados")
        filas = [[e["tipo"], fmin(e["inicio_s"]), fnum(e["duracion_s"], 1), fnum(e.get("v_max_kmh", e.get("pico")), 1),
                  fnum(e.get("distancia_m"), 0) if e["tipo"] == "Sprint" else f"{fnum(e.get('v_ini_kmh'), 1)} → {fnum(e.get('v_fin_kmh'), 1)}"]
                 for e in a["eventos"]]
        doc.tabla(["Tipo", "Inicio", "Duración (s)", "Pico (km/h o m/s²)", "Distancia (m) / velocidad (km/h)"], filas, hoja="Esfuerzos",
                  datos=[[e["tipo"], e["inicio_s"], e["duracion_s"], e.get("v_max_kmh", e.get("pico")), e.get("distancia_m")] for e in a["eventos"]])
    doc.salto()
    doc.h1("Calidad de la señal")
    q = a["calidad"]
    filas = []
    if q.get("gps"):
        g = q["gps"]
        filas += [["GPS", "Frecuencia de muestreo", f"{fnum(g['hz'], 0)} Hz"], ["GPS", "Muestras válidas", f"{fnum(g['validas_pct'], 1)} %"],
                  ["GPS", "Satélites (media; mínimo)", f"{fnum(g['sat_media'], 1)}; {fnum(g['sat_min'], 0)}"],
                  ["GPS", f"Muestras con < {cfg['sat_min']} satélites", f"{fnum(g['pct_sat_bajo'], 1)} %"],
                  ["GPS", "PDOP (media; P95)", f"{fnum(g['pdop_medio'], 2)}; {fnum(g['pdop_p95'], 2)}"],
                  ["GPS", f"Muestras con PDOP > {cfg['pdop_max']}", f"{fnum(g['pct_pdop_alto'], 1)} %"],
                  ["GPS", "Huecos > 0,5 s; > 2 s; máximo", f"{g['huecos_05s']}; {g['huecos_2s']}; {fnum(g['hueco_max_s'], 1)} s"],
                  ["GPS", "Velocidades imposibles descartadas", str(g["artefactos_v"])],
                  ["GPS", "Distancia ∫ velocidad / canal DIST / posiciones", f"{fnum(r.get('distancia_vel_m'), 0)} / {fnum(r.get('distancia_disp_m'), 0)} / {fnum(g.get('dist_haversine'), 0)} m"]]
    if q.get("acc"):
        g = q["acc"]
        filas += [["Acelerómetro", "Frecuencia de muestreo", f"{fnum(g['hz'], 0)} Hz"], ["Acelerómetro", "Huecos > 0,1 s", str(g["huecos_01s"])],
                  ["Acelerómetro", "Muestras saturadas (±16 g)", str(g["saturadas"])], ["Acelerómetro", "Mediana de |a| (≈ 1 g en reposo)", f"{fnum(g['mod_mediana'], 3)} g"]]
    if q.get("hr"):
        g = q["hr"]
        filas += [["Pulsómetro", "Cobertura del registro", f"{fnum(g['cobertura_pct'], 1)} %"], ["Pulsómetro", "Huecos > 5 s; máximo", f"{g['huecos_5s']}; {fnum(g['hueco_max_s'], 0)} s"],
                  ["Pulsómetro", "Valores descartados", str(g["descartadas"])]]
    doc.tabla(["Señal", "Control", "Resultado"], filas, hoja="Calidad")
    if r.get("distancia_vel_m") and r.get("distancia_disp_m"):
        dif = 100 * (r["distancia_disp_m"] - r["distancia_vel_m"]) / r["distancia_vel_m"]
        doc.nota(f"El canal DIST del dispositivo da un {fnum(dif, 1)} % más de distancia que la integral de la velocidad Doppler: DIST suma también "
                 "el «ruido» de la posición GPS cuando el jugador está quieto. Por eso, por defecto, se usa la velocidad Doppler (más fiable a baja velocidad).")
    doc.h2("Archivo y canales")
    doc.tabla(["Dato", "Valor"], [["Archivo", s["archivo"]], ["Formato", meta.get("formato", "")], ["Dispositivo", f"{meta.get('dispositivo', '')} (ID {meta.get('id', '—')})"],
                                  ["Firmware", meta.get("firmware", "—")], ["Jugador en el archivo", f"{meta.get('jugador', '—')} · dorsal {meta.get('dorsal', '—')}"],
                                  ["Inicio / fin", f"{ffecha(meta.get('inicio'))} – {ffecha(meta.get('fin'))}"], ["Zona horaria", f"{meta.get('zona', '')} {meta.get('desfase', '')}"],
                                  ["Paquetes leídos", fnum(meta.get("paquetes"), 0)], ["Bytes no reconocidos (resincronizaciones)", str(meta.get("resincronizaciones", 0))]])
    inv = a["inventario"]
    doc.tabla(["Sensor", "Paquete", "Magnitud", "Canales", "Muestras", "Hz"],
              [[i["sensor"], i["paquete"], i.get("magnitud") or "", i["canales"], fnum(i["muestras"], 0), fnum(i.get("hz"), 1)] for i in inv],
              hoja="Canales", datos=[[i["sensor"], i["paquete"], i.get("magnitud"), i["canales"], i["muestras"], i.get("hz")] for i in inv])
    doc.h1("Procesado")
    doc.lista(_parametros(cfg))
    return doc, a


# =========================================================================== colectivo y jugador
def _df_sel(opciones):
    df = SV.tabla(opciones.get("filtro") or {})
    if df.empty:
        raise ValueError("No hay sesiones con esos filtros")
    return df


def _ffila(row, ks):
    return [fvar(k, row.get(k)) for k in ks]


def informe_colectivo(opciones, carpeta_fig: Path):
    df = _df_sel(opciones).sort_values(["fecha", "jugador"])
    ks = [k for k in TABLA_JUG if k in df.columns and df[k].notna().any()]
    fechas = sorted(df["fecha"].dropna().unique())
    sub = (ffecha(fechas[0], False) if len(fechas) == 1 else f"{ffecha(fechas[0], False)} – {ffecha(fechas[-1], False)}") if fechas else ""
    tipos = ", ".join(sorted(set(df["tipo"]) - {""}))
    doc = Doc(opciones.get("titulo") or "Informe colectivo de carga", " · ".join(x for x in (sub, tipos) if x),
              [f"{len(df)} sesiones", f"{df['jugador_id'].nunique()} jugadores"])
    # una fila por jugador (media de sus sesiones si tiene varias)
    g = df.groupby("jugador")[ks].mean(numeric_only=True)
    n = df.groupby("jugador").size()
    doc.h1("Resumen del grupo")
    med = {k: df[k].mean() for k in ks}
    doc.kpis([(f"{KPI_LABEL.get(k, (M.etiqueta(k), ''))[0]} (media)", fvar(k, med[k]), KPI_LABEL.get(k, ('', ''))[1])
              for k in ("distancia_m", "m_min", "v_max_kmh", "hsr_m", "n_acel", "n_decel", "player_load", "fc_media") if k in med])
    doc.h1("Valores por jugador")
    desc = {k: E.describir(df[k].to_numpy(float)) for k in ks}
    bloques = [ks[i:i + 8] for i in range(0, len(ks), 8)]
    for bi, kb in enumerate(bloques):
        cols = ["Jugador", "Sesiones"] + [M.etiqueta(k) for k in kb]
        filas = [[j, str(int(n[j]))] + [fvar(k, row[k]) for k in kb] for j, row in g.iterrows()]
        filas.append(["Media ± DE", ""] + [f"{fvar(k, desc[k]['media'])} ± {fvar(k, desc[k]['de'])}" for k in kb])
        filas.append(["CV (%)", ""] + [fnum(desc[k]["cv_pct"], 1) for k in kb])
        ultimo = bi == len(bloques) - 1
        doc.tabla(cols, filas, nota="Si un jugador tiene varias sesiones se muestra la media. DE: desviación típica; CV: coeficiente de variación." if ultimo else None,
                  hoja="Por jugador" if bi == 0 else None, datos=[[j, int(n[j])] + [row[k] for k in ks] for j, row in g.iterrows()] if bi == 0 else None)
    if len(bloques) > 1:     # en Excel, todas las columnas en una hoja
        doc.b[-len(bloques)][1]["cols"] = ["Jugador", "Sesiones"] + [M.etiqueta(k) for k in ks]
    # destacados (± 1,5 DE)
    dest = []
    for k in ("distancia_m", "hsr_m", "n_acel", "n_decel", "player_load"):
        if k not in g or g[k].std() == 0 or not np.isfinite(g[k].std()):
            continue
        z = (g[k] - g[k].mean()) / g[k].std()
        for j, zz in z.items():
            if abs(zz) >= 1.5:
                dest.append(f"{j}: {M.etiqueta(k).lower()} {'muy por encima' if zz > 0 else 'muy por debajo'} del grupo ({fvar(k, g.loc[j, k])}; z = {fnum(zz, 1)}).")
    if dest:
        doc.h2("Valores destacados (más de 1,5 DE de la media del grupo)")
        doc.lista(dest)
    pref = carpeta_fig / f"col_{datetime.now():%H%M%S}"
    for k in ("distancia_m", "m_min", "hsr_m", "player_load", "fc_media"):
        if k in g and g[k].notna().any():
            v = g[k].dropna()
            doc.img(F.barras_jugadores(v.index.tolist(), v.to_numpy(float).tolist(), f"{pref}_{k}.png", M.VAR[k][1], M.VAR[k][2], float(v.mean())), ancho=.8)
    doc.h1("Sesiones incluidas")
    doc.tabla(["Fecha", "Jugador", "Tipo", "Tarea", "Archivo"], [[ffecha(r.inicio), r.jugador, r.tipo, r.tarea, r.archivo] for r in df.itertuples()],
              hoja="Sesiones")
    return doc, df


def informe_jugador(opciones, carpeta_fig: Path):
    df = _df_sel(opciones).sort_values("inicio")
    jug = df["jugador"].iloc[0]
    if df["jugador_id"].nunique() > 1:
        raise ValueError("Elija un único jugador para el informe de evolución")
    ks = [k for k in TABLA_JUG if k in df.columns and df[k].notna().any()]
    doc = Doc(f"Evolución de la carga · {jug}", f"{ffecha(df['inicio'].iloc[0], False)} – {ffecha(df['inicio'].iloc[-1], False)}",
              [f"{len(df)} sesiones"])
    doc.h1("Resumen")
    doc.kpis([(f"{KPI_LABEL.get(k, (M.etiqueta(k), ''))[0]} (media)", fvar(k, df[k].mean()), KPI_LABEL.get(k, ('', ''))[1])
              for k in ("distancia_m", "m_min", "v_max_kmh", "hsr_m", "n_acel", "player_load", "fc_media", "trimp") if k in df])
    if len(df) >= 3:
        ult = df.iloc[-1]
        prev = df.iloc[:-1]
        frases = []
        for k in ("distancia_m", "m_min", "hsr_m", "player_load", "fc_media"):
            if k in df and prev[k].notna().sum() >= 2 and prev[k].std() > 0 and pd.notna(ult[k]):
                z = (ult[k] - prev[k].mean()) / prev[k].std()
                frases.append(f"{M.VAR[k][1]}: {fvar(k, ult[k])} {M.VAR[k][2]} en la última sesión frente a {fvar(k, prev[k].mean())} de media "
                              f"(z = {fnum(z, 1)}{'; valor inusual' if abs(z) >= 2 else ''}).")
        if frases:
            doc.h2("Última sesión frente a las anteriores")
            doc.lista(frases)
    pref = carpeta_fig / f"jug_{datetime.now():%H%M%S}"
    fechas = [ffecha(x, False) for x in df["inicio"]]
    for k in ("distancia_m", "m_min", "hsr_m", "player_load", "fc_media"):
        if k in df and df[k].notna().sum() >= 2:
            doc.img(F.evolucion(fechas, df[k].to_numpy(float), f"{pref}_{k}.png", M.VAR[k][1], M.VAR[k][2]),
                    "Línea discontinua: media; banda: ± 1 DE.", .85)
    doc.h1("Sesiones")
    cols = ["Fecha", "Tipo"] + [M.etiqueta(k) for k in ks]
    doc.tabla(cols, [[ffecha(r["inicio"]), r["tipo"]] + _ffila(r, ks) for _, r in df.iterrows()], hoja="Sesiones",
              datos=[[r["inicio"], r["tipo"]] + [r[k] for k in ks] for _, r in df.iterrows()])
    return doc, df


# =========================================================================== informe estadístico (técnico colectivo)
def informe_estadistico(opciones, carpeta_fig: Path):
    cfg = SV.cfg()
    df = _df_sel(opciones)
    alfa = M._num(cfg.get("alfa"), .05)
    variables = opciones.get("variables") or ["distancia_m", "m_min", "v_max_kmh", "hsr_m", "n_acel", "n_decel", "player_load", "fc_media", "trimp"]
    variables = [v for v in variables if v in df.columns]
    doc = Doc(opciones.get("titulo") or "Informe técnico: análisis estadístico de la carga", cfg.get("titulo", ""),
              [f"{len(df)} sesiones", f"{df['jugador_id'].nunique()} jugadores", f"α = {fnum(alfa, 2)}"])
    pref = carpeta_fig / f"est_{datetime.now():%H%M%S}"
    doc.h1("Métodos")
    doc.p("Los datos se leyeron directamente de los archivos del dispositivo (.qul / .qui) con WIMU Lab. Procesado de las señales:")
    doc.lista(_parametros(cfg))
    doc.p("Estadística: descriptivos (media, DE, IC 95 %, mediana [P25–P75], CV y prueba de Shapiro–Wilk). Al aportar cada jugador varias "
          "sesiones, las comparaciones entre niveles independientes se hacen con la media de cada jugador (ANOVA de Welch o Kruskal–Wallis "
          "según la normalidad; comparaciones por parejas con t de Welch o U de Mann–Whitney, g de Hedges y corrección de Holm); las "
          "comparaciones de dos condiciones en los mismos jugadores con prueba t pareada o Wilcoxon y g de Hedges (av) con IC bootstrap; y "
          "el análisis con todas las sesiones mediante modelos lineales mixtos con intercepto aleatorio por jugador (REML). Correlaciones de "
          "Spearman y de medidas repetidas (rmcorr). Variabilidad entre sesiones: ICC(1) y CV intraindividual. Concordancia con SPRO: sesgo, "
          "MAE, RMSE, MAPE, ICC(A,1) e ICC(C,1) (McGraw y Wong, 1996; Koo y Li, 2016), CCC de Lin y Bland–Altman.")
    agr = opciones.get("agrupar") or ""
    d = A.descriptivo(df, variables, agr or None, alfa)
    doc.h1("Descriptivo" + (f" por {agr}" if agr else ""))
    filas = [[x["nombre"], x["grupo"], str(x["n"]), f"{fnum(x['media'], 1)} ± {fnum(x['de'], 1)}", f"[{fnum(x['ic_inf'], 1)}; {fnum(x['ic_sup'], 1)}]",
              f"{fnum(x['mediana'], 1)} [{fnum(x['p25'], 1)}–{fnum(x['p75'], 1)}]", fnum(x["cv_pct"], 1), E.fmt_p(x["sw_p"] if x["sw_p"] is not None else np.nan)]
             for x in d["filas"]]
    doc.tabla(["Variable", "Grupo", "n", "Media ± DE", "IC 95 %", "Mediana [P25–P75]", "CV (%)", "Shapiro p"], filas, hoja="Descriptivo",
              datos=[[x["nombre"], x["grupo"], x["n"], x["media"], x["de"], x["ic_inf"], x["ic_sup"], x["mediana"], x["p25"], x["p75"], x["cv_pct"], x["sw_p"]]
                     for x in d["filas"]])
    if agr:
        for v in variables[:6]:
            pts = d["puntos"].get(v) or []
            doc.img(F.puntos_grupos([(p["grupo"], p["valores"]) for p in pts], f"{pref}_d_{v}.png", M.VAR.get(v, (v, v))[1], M.VAR.get(v, ("", "", ""))[2]), ancho=.75)
    comp = opciones.get("comparacion")
    if comp and comp.get("factor"):
        try:
            c = A.comparar(df, variables, comp["factor"], comp.get("diseno", "grupos"), comp.get("niveles"), alfa)
            doc.h1(f"Comparación por {comp['factor']} ({ {'grupos': 'niveles independientes', 'pareado': 'medidas pareadas', 'mixto': 'modelo mixto'}.get(c['diseno'], c['diseno'])})")
            _tabla_comparacion(doc, c)
        except Exception as e:  # noqa: BLE001
            doc.nota(f"No se pudo hacer la comparación: {e}")
    doc.h1("Correlaciones entre variables (Spearman)")
    mt = A.matriz(df, variables)
    doc.tabla([""] + mt["nombres"], [[mt["nombres"][i]] + [fnum(c["rho"], 2) + ("*" if c["p"] is not None and c["p"] < alfa else "") for c in fila]
                                     for i, fila in enumerate(mt["m"])], nota=f"* p < {fnum(alfa, 2)} (sin corregir). Unidad: la sesión.", hoja="Correlaciones")
    if df["jugador_id"].nunique() >= 2 and df.groupby("jugador_id").size().max() >= 2:
        doc.h1("Variabilidad entre sesiones del mismo jugador")
        vb = A.variabilidad(df, variables, alfa)
        doc.tabla(["Variable", "n", "Jugadores", "ICC(1) [IC 95 %]", "Magnitud", "CV intraindividual (%)"],
                  [[x["nombre"], str(x["n"]), str(x["jugadores"]), f"{fnum(x['icc'], 2)} [{fnum(x['icc_ic'][0], 2)}; {fnum(x['icc_ic'][1], 2)}]", x["magnitud"],
                    fnum(x["cv_intra"], 1)] for x in vb], hoja="Variabilidad",
                  nota="ICC(1): proporción de la variabilidad debida a las diferencias entre jugadores; CV intraindividual: variación de cada jugador entre sesiones.")
    if opciones.get("incluir_spro", True):
        d_spro = SP.cargar(cfg["carpeta_datos"])
        if d_spro:
            con = SP.concordancia(d_spro, SV.reg().sesiones(incluir_faltan=False), {s["id"]: s.get("kpis") for s in SV.reg().sesiones()})
            if con["resultados"]:
                doc.salto()
                doc.h1("Concordancia con SPRO")
                _tabla_concordancia(doc, con, pref)
    return doc, df


def _tabla_comparacion(doc, c):
    res = c["resultados"]
    if c["diseno"] == "pareado":
        doc.tabla(["Variable", "n", c["niveles"][0], c["niveles"][1], "Diferencia [IC 95 %]", "Cambio %", "Prueba", "p", "p Holm", "Tamaño del efecto [IC 95 %]"],
                  [[r["nombre"], str(r.get("n", "")), f"{fnum(r.get('media_a'), 1)} ± {fnum(r.get('de_a'), 1)}", f"{fnum(r.get('media_b'), 1)} ± {fnum(r.get('de_b'), 1)}",
                    f"{fnum(r.get('dif'), 1)} [{fnum((r.get('dif_ic') or [None])[0], 1)}; {fnum((r.get('dif_ic') or [None, None])[1], 1)}]", fnum(r.get("cambio_pct"), 1),
                    r.get("prueba", r.get("error", "")), r.get("p_txt", "—"), r.get("p_holm_txt", "—"),
                    f"{r.get('te', '')} {fnum(r.get('te_valor'), 2)} [{fnum((r.get('te_ic') or [None])[0], 2)}; {fnum((r.get('te_ic') or [None, None])[1], 2)}] {r.get('magnitud', '')}"]
                   for r in res], hoja="Comparación", nota="Medias ± DE de la media de cada jugador en cada condición.")
    elif c["diseno"] == "mixto":
        filas = []
        for r in res:
            if r.get("error"):
                filas.append([r["nombre"], "—", r["error"], "", "", "", ""])
                continue
            for e in r["efectos"][1:]:
                filas.append([r["nombre"], e["termino"], f"{fnum(e['estimacion'], 2)} [{fnum(e['ic'][0], 2)}; {fnum(e['ic'][1], 2)}]", E.fmt_p(e["p"] if e["p"] is not None else np.nan),
                              fnum(e.get("d"), 2), fnum(r.get("icc"), 2), f"{fnum(r.get('r2m'), 2)} / {fnum(r.get('r2c'), 2)}"])
        doc.tabla(["Variable", "Efecto", "Estimación [IC 95 %]", "p", "d (DE total)", "ICC", "R² marg. / cond."], filas, hoja="Modelo mixto",
                  nota="Modelo lineal mixto con intercepto aleatorio por jugador (REML); grados de libertad between-within.")
    else:
        doc.tabla(["Variable", "Niveles (media ± DE; n)", "Prueba", "Estadístico", "p", "p Holm", "Tamaño del efecto"],
                  [[r["nombre"], "; ".join(f"{m['nivel']}: {fnum(m['media'], 1)} ± {fnum(m['de'], 1)} ({m['n']})" for m in r.get("medias", [])),
                    r.get("prueba", r.get("error", "")), fnum(r.get("estadistico"), 2), r.get("p_txt", "—"), r.get("p_holm_txt", "—"),
                    f"{r.get('te', '')} = {fnum(r.get('te_valor'), 3)}"] for r in res], hoja="Comparación",
                  nota="Unidad de análisis: la media de cada jugador en cada nivel.")
        pares = [(r["nombre"], p) for r in res for p in r.get("pares", [])]
        if pares:
            doc.h2("Comparaciones por parejas")
            doc.tabla(["Variable", "Pareja", "Diferencia", "Prueba", "p", "p Holm", "g de Hedges"],
                      [[n, f"{p['b']} − {p['a']}", fnum(p["dif"], 1), p["prueba"], E.fmt_p(p["p"] if p["p"] is not None else np.nan),
                        E.fmt_p(p["p_holm"] if p["p_holm"] is not None else np.nan), f"{fnum(p['g'], 2)} {p['magnitud']}"] for n, p in pares], hoja="Parejas")


def _tabla_concordancia(doc, con, pref):
    rs = [r for r in con["resultados"] if r.get("n", 0) >= 3]
    doc.p(f"Filas de SPRO: {con['n_filas']}; emparejadas con una sesión: {len(con['emparejadas'])}. Diferencia = WIMU Lab − SPRO.")
    doc.tabla(["Variable", "n", "SPRO", "WIMU Lab", "Sesgo [IC 95 %]", "Sesgo %", "MAE", "RMSE", "MAPE %", "ICC(A,1) [IC 95 %]", "Límites de acuerdo 95 %"],
              [[r["nombre"], str(r["n"]), f"{fnum(r['media_ref'], 1)} ± {fnum(r['de_ref'], 1)}", f"{fnum(r['media_nuevo'], 1)} ± {fnum(r['de_nuevo'], 1)}",
                f"{fnum(r['sesgo'], 2)} [{fnum(r['sesgo_ic'][0], 2)}; {fnum(r['sesgo_ic'][1], 2)}]", fnum(r["sesgo_pct"], 1), fnum(r["mae"], 2), fnum(r["rmse"], 2),
                fnum(r.get("mape"), 1), f"{fnum(r['icc_a'], 3)} [{fnum(r['icc_a_ic'][0], 3)}; {fnum(r['icc_a_ic'][1], 3)}] {r.get('magnitud', '')}",
                f"{fnum(r['lim_inf'], 2)} a {fnum(r['lim_sup'], 2)}"] for r in rs], hoja="Concordancia SPRO",
              datos=[[r["nombre"], r["n"], r["media_ref"], r["media_nuevo"], r["sesgo"], r["sesgo_pct"], r["mae"], r["rmse"], r.get("mape"), r["icc_a"], r["icc_c"],
                      r["ccc"], r.get("r"), r["lim_inf"], r["lim_sup"], r["prop_pendiente"], r["prop_p"]] for r in rs],
              nota="ICC(A,1): acuerdo absoluto; magnitud según Koo y Li (2016). MAPE solo con valores de referencia ≥ 1 % de la media.")
    for r in rs:
        doc.img(F.bland_altman(r, f"{pref}_ba_{r['variable']}.png"),
                f"Sesgo {fnum(r['sesgo'], 2)}; límites {fnum(r['lim_inf'], 2)} a {fnum(r['lim_sup'], 2)}; sesgo proporcional: pendiente {fnum(r['prop_pendiente'], 3)} (p {E.fmt_p(r['prop_p'])}).")
    if con["sin_pareja"]:
        doc.nota("Filas de SPRO sin sesión emparejada: " + "; ".join(f"fila {x['fila']} ({x['texto']})" for x in con["sin_pareja"][:30]))


# =========================================================================== generación
def generar(opciones, progreso=lambda f, m: None):
    """opciones: tipo, ids (sesiones) o filtro, formatos [docx, pdf, xlsx], …  Devuelve la lista de archivos."""
    cfg = SV.cfg()
    tipo = opciones.get("tipo", "sesion")
    formatos = opciones.get("formatos") or ["docx", "pdf", "xlsx"]
    dia = C.carpeta_dia(cfg)
    sub = {"sesion": "Sesiones", "tecnico": "Sesiones_tecnico", "jugador": "Jugadores", "colectivo": "Colectivo", "estadistico": "Estadistico"}[tipo]
    carpeta = dia / sub
    figs = dia / "_figuras"
    trabajos = []
    if tipo in ("sesion", "tecnico"):
        for sid in opciones.get("ids") or []:
            trabajos.append(("s", sid))
    elif tipo == "jugador" and opciones.get("jugadores"):
        for jid in opciones["jugadores"]:
            trabajos.append(("j", jid))
    else:
        trabajos.append(("g", None))
    if not trabajos:
        raise ValueError("No hay sesiones seleccionadas")
    salida = []
    for i, (k, x) in enumerate(trabajos):
        progreso(i / len(trabajos), f"Preparando informe {i + 1} de {len(trabajos)}…")
        if k == "s":
            doc, a = informe_sesion(x, figs, tecnico=(tipo == "tecnico"))
            s = a["sesion"]
            nombre = _slug(f"{(a['meta'].get('inicio') or '')[:10]}_{s.get('jugador_nombre') or s.get('dispositivo')}_{Path(s['archivo']).stem}")
            extra = {"Serie 1 Hz": pd.DataFrame({k2: v for k2, v in a["series"].items() if k2 != "mapa" and isinstance(v, list) and len(v) == len(a["series"]["t"])})} if tipo == "tecnico" else None
        elif k == "j":
            op = dict(opciones)
            op["filtro"] = dict(opciones.get("filtro") or {}, jugadores=[x])
            doc, df = informe_jugador(op, figs)
            nombre = _slug(f"Evolucion_{df['jugador'].iloc[0]}")
            extra = {"Datos": df}
        else:
            fn = {"colectivo": informe_colectivo, "estadistico": informe_estadistico, "jugador": informe_jugador}[tipo]
            doc, df = fn(opciones, figs)
            nombre = _slug(f"{'Colectivo' if tipo == 'colectivo' else 'Estadistico' if tipo == 'estadistico' else 'Evolucion'}_{datetime.now():%H%M%S}")
            extra = {"Datos": df}
        base = carpeta / nombre
        if "docx" in formatos:
            progreso((i + .3) / len(trabajos), f"Word: {nombre}")
            salida.append(str(a_docx(doc, base.with_suffix(".docx"))))
        if "pdf" in formatos or "html" in formatos:
            progreso((i + .55) / len(trabajos), f"PDF: {nombre}")
            h = a_html(doc, base.with_suffix(".html"))
            if "pdf" in formatos and html_a_pdf(h, base.with_suffix(".pdf")):
                salida.append(str(base.with_suffix(".pdf")))
                if "html" not in formatos:
                    h.unlink(missing_ok=True)
            else:
                salida.append(str(h))
        if "xlsx" in formatos:
            progreso((i + .85) / len(trabajos), f"Excel: {nombre}")
            salida.append(str(a_xlsx(doc, base.with_suffix(".xlsx"), extra)))
    return dict(archivos=salida, carpeta=str(carpeta))
